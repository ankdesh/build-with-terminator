"""Script to generate a synthetic sample XLSX file for testing xls2datasets.

Creates a multi-sheet workbook with:
  - Sheet "SensorData": Temperature/humidity sensor readings over time
  - Sheet "Experiments": Two separate tables (parameters + results)
  - Sheet "Summary": Summary statistics table
"""

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows


def generate_sample_xlsx(output_path: str = "sample_data.xlsx") -> Path:
    """Generate a synthetic XLSX file with multiple sheets and tables.

    Args:
        output_path: Path for the output XLSX file.

    Returns:
        Path to the generated file.
    """
    xlsx_file = Path(output_path)
    wb = Workbook()

    # --- Sheet 1: SensorData ---
    ws_sensor = wb.active
    ws_sensor.title = "SensorData"

    sensor_data = {
        "Timestamp": [
            "2026-01-15 08:00", "2026-01-15 08:15", "2026-01-15 08:30",
            "2026-01-15 08:45", "2026-01-15 09:00", "2026-01-15 09:15",
            "2026-01-15 09:30", "2026-01-15 09:45", "2026-01-15 10:00",
            "2026-01-15 10:15",
        ],
        "Sensor_ID": [
            "S001", "S001", "S001", "S001", "S001",
            "S002", "S002", "S002", "S002", "S002",
        ],
        "Temperature_C": [22.5, 22.8, 23.1, 23.4, 23.2, 21.0, 21.3, 21.8, 22.0, 22.1],
        "Humidity_Pct": [45.2, 44.8, 44.1, 43.5, 43.9, 52.1, 51.8, 51.2, 50.9, 50.5],
        "Battery_V": [3.72, 3.71, 3.71, 3.70, 3.70, 3.65, 3.64, 3.64, 3.63, 3.63],
        "Status": [
            "OK", "OK", "OK", "OK", "OK",
            "OK", "OK", "WARN", "OK", "OK",
        ],
    }
    df_sensor = pd.DataFrame(sensor_data)
    for row in dataframe_to_rows(df_sensor, index=False, header=True):
        ws_sensor.append(row)

    # --- Sheet 2: Experiments (two separate tables) ---
    ws_exp = wb.create_sheet("Experiments")

    # Table 1: Experiment Parameters (starts at row 1)
    ws_exp.append(["Experiment Parameters"])
    exp_params = {
        "Experiment_ID": ["EXP-001", "EXP-002", "EXP-003", "EXP-004"],
        "Material": ["Aluminum", "Steel", "Copper", "Titanium"],
        "Temperature_K": [300, 350, 400, 450],
        "Pressure_MPa": [1.0, 2.5, 5.0, 10.0],
        "Duration_Hours": [24, 48, 72, 96],
    }
    df_params = pd.DataFrame(exp_params)
    for row in dataframe_to_rows(df_params, index=False, header=True):
        ws_exp.append(row)

    # Leave some blank rows as separator
    ws_exp.append([])
    ws_exp.append([])

    # Table 2: Experiment Results (starts after blank rows)
    ws_exp.append(["Experiment Results"])
    exp_results = {
        "Experiment_ID": ["EXP-001", "EXP-002", "EXP-003", "EXP-004"],
        "Yield_Strength_MPa": [276.0, 520.0, 210.0, 880.0],
        "Elongation_Pct": [12.5, 8.2, 15.0, 6.1],
        "Hardness_HV": [75, 180, 55, 310],
        "Pass_Fail": ["PASS", "PASS", "FAIL", "PASS"],
    }
    df_results = pd.DataFrame(exp_results)
    for row in dataframe_to_rows(df_results, index=False, header=True):
        ws_exp.append(row)

    # --- Sheet 3: Summary ---
    ws_summary = wb.create_sheet("Summary")

    summary_data = {
        "Metric": [
            "Total Experiments", "Pass Rate", "Avg Yield Strength",
            "Avg Elongation", "Max Hardness",
        ],
        "Value": [4, 0.75, 471.5, 10.45, 310],
        "Unit": ["count", "fraction", "MPa", "%", "HV"],
    }
    df_summary = pd.DataFrame(summary_data)
    for row in dataframe_to_rows(df_summary, index=False, header=True):
        ws_summary.append(row)

    wb.save(str(xlsx_file))
    print(f"Generated sample XLSX file: {xlsx_file.resolve()}")
    return xlsx_file


if __name__ == "__main__":
    out_path = sys.argv[1] if len(sys.argv) > 1 else "sample_data.xlsx"
    generate_sample_xlsx(out_path)
